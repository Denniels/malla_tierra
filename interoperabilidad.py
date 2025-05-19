# -*- coding: utf-8 -*-
"""
Módulo para manejo de interoperabilidad con otros sistemas
"""
import json
import csv
from typing import Dict, List, Any, Optional
import xml.etree.ElementTree as ET
from dataclasses import dataclass, asdict
from datetime import datetime
from pathlib import Path
import numpy as np

import ezdxf
from fastapi import FastAPI, HTTPException
import openpyxl
from pydantic import BaseModel

from malla import MallaTierra, Conductor, Varilla
from analisis_avanzado import CapaSuelo

@dataclass
class DatosExportacion:
    """Estructura de datos para exportación"""
    tipo: str
    datos: Dict[str, Any]
    metadatos: Dict[str, str]
    fecha: str

class DatosMedicion(BaseModel):
    """Modelo para datos de medición de resistividad"""
    puntos: List[Dict[str, float]]  # Lista de puntos medidos {x, y, resistividad}
    fecha: str
    observaciones: Optional[str] = None

class ExportadorCAD:
    """Clase para exportar diseños a formatos CAD"""
    
    @staticmethod
    def exportar_dxf(malla: MallaTierra, ruta: str) -> bool:
        """
        Exporta la malla a formato DXF usando ezdxf.
        Compatible con AutoCAD y software similar.
        """
        try:
            doc = ezdxf.new("R2010")  # Versión AutoCAD 2010
            msp = doc.modelspace()
            
            # Crear capa para conductores
            doc.layers.new("CONDUCTORES", dxfattribs={'color': 1})  # 1 = rojo
            
            # Crear capa para varillas
            doc.layers.new("VARILLAS", dxfattribs={'color': 5})  # 5 = azul
            
            # Conductores horizontales
            for i in range(malla.n_x):
                y = np.linspace(0, malla.largo, malla.n_y)
                for j in range(len(y)):
                    msp.add_line(
                        (i * malla.espaciamiento_x, y[j], -malla.profundidad),
                        ((i + 1) * malla.espaciamiento_x, y[j], -malla.profundidad),
                        dxfattribs={'layer': 'CONDUCTORES'}
                    )
            
            # Conductores verticales
            for i in range(malla.n_y):
                x = np.linspace(0, malla.ancho, malla.n_x)
                for j in range(len(x)):
                    msp.add_line(
                        (x[j], i * malla.espaciamiento_y, -malla.profundidad),
                        (x[j], (i + 1) * malla.espaciamiento_y, -malla.profundidad),
                        dxfattribs={'layer': 'CONDUCTORES'}
                    )
            
            # Varillas
            for varilla in malla.varillas:
                msp.add_line(
                    (varilla.posicion_x, varilla.posicion_y, -malla.profundidad),
                    (varilla.posicion_x, varilla.posicion_y, -(malla.profundidad + varilla.longitud)),
                    dxfattribs={'layer': 'VARILLAS'}
                )
            
            # Guardar archivo
            doc.saveas(ruta)
            return True
            
        except Exception as e:
            print(f"Error al exportar a DXF: {str(e)}")
            return False

class ImportadorMediciones:
    """Clase para importar datos de mediciones de campo"""
    
    @staticmethod
    def importar_csv(ruta: str) -> DatosMedicion:
        """
        Importa mediciones desde un archivo CSV.
        Formato esperado:
        x,y,resistividad,observaciones
        """
        datos = {
            'puntos': [],
            'fecha': datetime.now().isoformat(),
            'observaciones': []
        }
        
        try:
            with open(ruta, 'r') as f:
                reader = csv.DictReader(f)
                for row in reader:
                    datos['puntos'].append({
                        'x': float(row['x']),
                        'y': float(row['y']),
                        'resistividad': float(row['resistividad'])
                    })
                    if row.get('observaciones'):
                        datos['observaciones'].append(row['observaciones'])
            
            if datos['observaciones']:
                datos['observaciones'] = "\n".join(datos['observaciones'])
            else:
                datos['observaciones'] = None
                
            return DatosMedicion(**datos)
            
        except Exception as e:
            print(f"Error al importar CSV: {str(e)}")
            return DatosMedicion(puntos=[], fecha=datetime.now().isoformat())
    
    @staticmethod
    def importar_excel(ruta: str) -> DatosMedicion:
        """
        Importa mediciones desde un archivo Excel.
        Formato esperado:
        Hoja 1:
        A1: Fecha
        A2: Observaciones
        A4+: x, y, resistividad
        """
        try:
            wb = openpyxl.load_workbook(ruta)
            ws = wb.active
            
            datos = {
                'puntos': [],
                'fecha': ws['A1'].value or datetime.now().isoformat(),
                'observaciones': ws['A2'].value
            }
            
            # Leer puntos de medición (asume formato x,y,resistividad)
            for row in ws.iter_rows(min_row=4):  # Empieza en fila 4
                if not row[0].value:  # Si llegamos a una fila vacía
                    break
                datos['puntos'].append({
                    'x': float(row[0].value),
                    'y': float(row[1].value),
                    'resistividad': float(row[2].value)
                })
                
            return DatosMedicion(**datos)
            
        except Exception as e:
            print(f"Error al importar Excel: {str(e)}")
            return DatosMedicion(puntos=[], fecha=datetime.now().isoformat())

class IntegracionSoftware:
    """Clase para integración con software de diseño eléctrico"""
    
    @staticmethod
    def exportar_etap(malla: MallaTierra, ruta: str) -> bool:
        """
        Exporta datos en formato compatible con ETAP.
        """
        try:
            root = ET.Element("MallaTierra")
            
            # Datos generales
            datos = ET.SubElement(root, "DatosGenerales")
            ET.SubElement(datos, "Ancho").text = str(malla.ancho)
            ET.SubElement(datos, "Largo").text = str(malla.largo)
            ET.SubElement(datos, "Profundidad").text = str(malla.profundidad)
            
            # Conductores
            conductores = ET.SubElement(root, "Conductores")
            ET.SubElement(conductores, "Material").text = malla.conductor.material
            ET.SubElement(conductores, "Diametro").text = str(malla.conductor.diametro)
            
            # Varillas
            varillas = ET.SubElement(root, "Varillas")
            for v in malla.varillas:
                varilla = ET.SubElement(varillas, "Varilla")
                ET.SubElement(varilla, "X").text = str(v.posicion_x)
                ET.SubElement(varilla, "Y").text = str(v.posicion_y)
                ET.SubElement(varilla, "Longitud").text = str(v.longitud)
            
            # Guardar archivo
            tree = ET.ElementTree(root)
            tree.write(ruta, encoding='utf-8', xml_declaration=True)
            
            return True
        except Exception as e:
            print(f"Error al exportar a ETAP: {str(e)}")
            return False

    @staticmethod
    def exportar_excel(malla: MallaTierra, parametros: Dict, resultados: Dict, ruta: str) -> bool:
        """
        Exporta los resultados a formato Excel
        """
        try:
            wb = openpyxl.Workbook()
            
            # Hoja de parámetros
            ws_params = wb.active
            ws_params.title = "Parámetros"
            ws_params.append(["Parámetro", "Valor", "Unidad"])
            
            unidades = {
                'resistividad': 'Ω⋅m',
                'corriente': 'A',
                'tiempo': 's',
                'longitud': 'm',
                'area': 'm²',
                'temperatura': '°C',
                'resistencia': 'Ω',
                'voltaje': 'V',
                'potencial': 'V'
            }
            
            # Agregar parámetros principales
            for param, valor in parametros.items():
                ws_params.append([param, valor, unidades.get(param, '')])
                
            # Hoja de resultados
            ws_res = wb.create_sheet("Resultados")
            ws_res.append(["Variable", "Valor", "Unidad"])
            
            for var, valor in resultados.items():
                ws_res.append([var, valor, unidades.get(var, '')])
                
            # Guardar archivo
            wb.save(ruta)
            return True
            
        except Exception as e:
            print(f"Error al exportar a Excel: {str(e)}")
            return False

class APIIntegracion:
    """Maneja la serialización y deserialización para la API"""
    
    @staticmethod
    def generar_json(malla: MallaTierra) -> str:
        """
        Genera una representación JSON de la malla.
        """
        datos = {
            'tipo': 'malla_tierra',
            'version': '1.0',
            'parametros': {
                'ancho': malla.ancho,
                'largo': malla.largo,
                'profundidad': malla.profundidad,
                'espaciamiento_x': malla.espaciamiento_x,
                'espaciamiento_y': malla.espaciamiento_y,
                'conductor': {
                    'material': malla.conductor.material,
                    'diametro': malla.conductor.diametro
                },
                'varillas': [
                    {
                        'x': v.posicion_x,
                        'y': v.posicion_y,
                        'longitud': v.longitud,
                        'diametro': v.diametro
                    }
                    for v in malla.varillas
                ]
            }
        }
        
        return json.dumps(datos, indent=2)
    
    @staticmethod
    def cargar_json(datos_json: str) -> Optional[MallaTierra]:
        """
        Crea una malla a partir de datos JSON.
        """
        try:
            datos = json.loads(datos_json)
            
            # Crear conductor
            conductor_datos = datos['parametros']['conductor']
            if conductor_datos['material'] == 'Cobre':
                conductor = Conductor.get_conductor_cobre(conductor_datos['diametro'])
            else:
                conductor = Conductor.get_conductor_acero(conductor_datos['diametro'])
            
            # Crear malla
            malla = MallaTierra(
                ancho=datos['parametros']['ancho'],
                largo=datos['parametros']['largo'],
                espaciamiento=datos['parametros']['espaciamiento_x'],
                conductor=conductor,
                profundidad=datos['parametros']['profundidad']
            )
            
            # Agregar varillas
            for v in datos['parametros']['varillas']:
                malla.agregar_varilla(
                    x=v['x'],
                    y=v['y'],
                    longitud=v['longitud'],
                    diametro=v['diametro']
                )
            
            return malla
        except Exception as e:
            print(f"Error al cargar JSON: {str(e)}")
            return None

# API REST para integración
app = FastAPI(
    title="API Malla de Tierra",
    description="API REST para integración con el sistema de malla de tierra",
    version="1.0.0"
)

# Variables globales para la API
malla_activa: Optional[MallaTierra] = None
parametros_actuales: Dict[str, Any] = {}
resultados_actuales: Dict[str, Any] = {}

@app.get("/parametros")
async def obtener_parametros() -> Dict[str, Any]:
    """Obtiene los parámetros configurados"""
    return parametros_actuales

@app.post("/mediciones")
async def cargar_mediciones(datos: DatosMedicion) -> Dict[str, Any]:
    """Carga nuevos datos de medición"""
    try:
        # Aquí se procesarían los datos y se actualizaría el modelo
        return {
            "mensaje": "Datos de medición cargados correctamente",
            "puntos_procesados": len(datos.puntos)
        }
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

@app.get("/resultados")
async def obtener_resultados() -> Dict[str, Any]:
    """Obtiene los resultados del último cálculo"""
    if not resultados_actuales:
        raise HTTPException(
            status_code=404,
            detail="No hay resultados disponibles. Debe realizar un cálculo primero."
        )
    return resultados_actuales

@app.post("/exportar")
async def exportar_diseño(formato: str, ruta: str) -> Dict[str, Any]:
    """
    Exporta el diseño actual en el formato especificado
    
    Formatos soportados:
    - dxf: AutoCAD DXF
    - etap: ETAP XML
    - excel: Excel XLSX
    - json: JSON
    """
    if not malla_activa:
        raise HTTPException(
            status_code=404,
            detail="No hay un diseño activo para exportar"
        )
    
    try:
        ruta_base = Path(ruta)
        if formato == "dxf":
            ruta_salida = str(ruta_base.with_suffix(".dxf"))
            if ExportadorCAD.exportar_dxf(malla_activa, ruta_salida):
                return {"mensaje": f"Diseño exportado a {ruta_salida}"}
        
        elif formato == "etap":
            ruta_salida = str(ruta_base.with_suffix(".xml"))
            if IntegracionSoftware.exportar_etap(malla_activa, ruta_salida):
                return {"mensaje": f"Diseño exportado a {ruta_salida}"}
        
        elif formato == "excel":
            ruta_salida = str(ruta_base.with_suffix(".xlsx"))
            if IntegracionSoftware.exportar_excel(
                malla_activa,
                parametros_actuales,
                resultados_actuales,
                ruta_salida
            ):
                return {"mensaje": f"Diseño exportado a {ruta_salida}"}
        
        elif formato == "json":
            ruta_salida = str(ruta_base.with_suffix(".json"))
            datos = APIIntegracion.generar_json(malla_activa)
            with open(ruta_salida, 'w', encoding='utf-8') as f:
                f.write(datos)
            return {"mensaje": f"Diseño exportado a {ruta_salida}"}
        
        else:
            raise HTTPException(
                status_code=400,
                detail=f"Formato no soportado: {formato}"
            )
            
    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"Error al exportar: {str(e)}"
        )
